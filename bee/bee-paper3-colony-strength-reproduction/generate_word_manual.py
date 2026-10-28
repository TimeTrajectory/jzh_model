"""Generate the detailed VS Code training and project structure manual as DOCX."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "docs/第三篇论文复现项目目录结构与VSCode训练操作手册.docx"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
DARK = RGBColor(31, 78, 120)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        set_cell_shading(table.rows[0].cells[index], LIGHT_BLUE)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_text(cells[index], value)
            if row_index % 2:
                set_cell_shading(cells[index], "F8FBFD")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Code"]
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.right_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    run = paragraph.add_run(text.rstrip())
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)


def add_bullets(doc: Document, items: list[str], level: int = 0) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        paragraph.add_run(item)


def add_steps(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_note(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF2CC")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(f"{title}：")
    run.bold = True
    paragraph.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_toc(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "在 Word 中右键此处，选择“更新域/更新整个目录”生成页码目录。"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    paragraph.add_run(" 页")


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 24), ("Subtitle", 13), ("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = DARK
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 3"].paragraph_format.space_before = Pt(7)
    code = doc.styles.add_style("Code", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(8.5)

    header = section.header.paragraphs[0]
    header.text = "第三篇论文复现项目 · VS Code 目录结构与训练操作手册"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(section.footer.paragraphs[0])
    return doc


def build_manual() -> Document:
    doc = setup_document()
    props = doc.core_properties
    props.title = "第三篇论文复现项目目录结构与 VS Code 训练操作手册"
    props.subject = "YOLOv10X 蜜蜂三类别检测项目的目录、训练、权重、评估与修改指南"
    props.author = "TimeTrajectory 蜜蜂论文复现项目"
    props.keywords = "YOLOv10X, VS Code, 蜜蜂检测, 训练, 论文复现"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("第三篇论文复现项目\n目录结构与 VS Code 训练操作手册")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("YOLOv10X 三类别蜜蜂检测 · 从数据准备到权重管理与二次开发")
    doc.add_paragraph()
    cover = add_table(doc, ["项目", "内容"], [
        ["项目目录", "/jzk-model/jzh_model/bee/bee-paper3-colony-strength-reproduction"],
        ["对应论文", "Chaudhary et al., Smart Agricultural Technology 14 (2026) 102005"],
        ["检测类别", "WB no pollen、WB pollen、Drone bee"],
        ["主要框架", "Ultralytics 8.4.163 + YOLOv10X + PyTorch"],
        ["适用编辑器", "VS Code，本机 Linux、Remote SSH 或 WSL 均可"],
        ["手册版本", f"1.0（{date.today().isoformat()}）"],
    ], [4.0, 12.5])
    doc.add_paragraph()
    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.add_run("用途：理解项目、独立训练、修改参数、替换权重、扩展数据与排查常见问题").bold = True
    doc.add_page_break()

    doc.add_heading("目录", level=1)
    add_toc(doc)
    add_note(doc, "目录更新", "第一次打开 DOCX 时，Word/WPS 可能只显示目录提示文字。右键目录并选择“更新整个目录”，即可生成章节页码。")

    doc.add_heading("1. 项目定位与整体工作流", level=1)
    doc.add_paragraph("本项目复现第三篇论文中的 YOLOv10X 实时目标检测路线。输入是蜂箱入口图片或视频帧，输出是三类蜜蜂边界框、类别和置信度。训练代码、数据转换、测试评估、推理和报告生成均放在同一个项目目录内。")
    doc.add_heading("1.1 三个检测类别", level=2)
    add_table(doc, ["类别 ID", "英文名称", "中文含义", "当前原始标注数"], [
        [0, "WB no pollen", "未携带花粉的工蜂", "73,981"],
        [1, "WB pollen", "携带花粉的工蜂", "4,272"],
        [2, "Drone bee", "雄蜂", "559"],
    ], [2.0, 4.0, 5.0, 3.0])
    doc.add_heading("1.2 从原始数据到预测结果", level=2)
    add_steps(doc, [
        "下载官方 supplement.zip，并放入 downloads/。",
        "运行 inspect_archive.py 检查压缩包结构和哈希。",
        "运行 prepare_dataset.py，把 LabelMe JSON 转成 YOLO 标签，并生成固定的 train/val/test 划分。",
        "运行 train_yolov10.py，从 COCO 预训练权重开始训练；best.pt 和 last.pt 写入 runs/。",
        "运行 evaluate.py，在固定测试集计算整体和逐类别指标。",
        "运行 predict.py，对自己的图片或视频进行推理。",
        "运行 generate_report.py 或 generate_word_manual.py 生成报告与操作文档。",
    ])
    add_code(doc, "官方 ZIP → inspect_archive.py → prepare_dataset.py → dataset.yaml\n"
                  "dataset.yaml + 初始权重 → train_yolov10.py → best.pt / last.pt\n"
                  "best.pt → evaluate.py / predict.py → 指标、曲线、预测图片")

    doc.add_heading("2. 在 VS Code 中正确打开项目", level=1)
    doc.add_heading("2.1 推荐打开方式", level=2)
    doc.add_paragraph("在 VS Code 中选择“文件 → 打开文件夹”，直接打开下面这一层目录，而不是只打开某个 Python 文件：")
    add_code(doc, "/jzk-model/jzh_model/bee/bee-paper3-colony-strength-reproduction")
    doc.add_paragraph("打开正确后，VS Code 左侧资源管理器顶层应直接看到 README.md、train_yolov10.py、evaluate.py、data、runs、weights 等内容。这样相对路径、终端工作目录和 Python 导入最不容易出错。")
    doc.add_heading("2.2 建议安装的 VS Code 扩展", level=2)
    add_table(doc, ["扩展", "用途", "是否必须"], [
        ["Python（Microsoft）", "运行、调试、选择解释器、测试 Python 文件", "建议安装"],
        ["Pylance", "代码补全、类型提示、跳转定义", "建议安装"],
        ["Remote - SSH", "VS Code 连接远程 GPU 服务器", "远程训练时安装"],
        ["WSL", "在 Windows 的 WSL Linux 环境中打开项目", "使用 WSL 时安装"],
        ["Git Graph / GitLens", "查看实验代码改动和提交历史", "可选"],
    ], [4.0, 8.0, 3.0])
    doc.add_heading("2.3 选择项目虚拟环境", level=2)
    add_steps(doc, [
        "按 Ctrl+Shift+P 打开命令面板。",
        "输入并选择 Python: Select Interpreter。",
        "选择项目中的 .venv/bin/python。",
        "如果列表中没有，选择 Enter interpreter path，手动输入完整路径。",
    ])
    add_code(doc, "/jzk-model/jzh_model/bee/bee-paper3-colony-strength-reproduction/.venv/bin/python")
    doc.add_paragraph("VS Code 右下角应显示该解释器。后续点击“运行 Python 文件”时，才会使用已经安装 Ultralytics、PyTorch 和 OpenCV 的环境。")
    doc.add_heading("2.4 打开集成终端并检查位置", level=2)
    add_code(doc, "pwd\nls\n.venv/bin/python --version")
    doc.add_paragraph("pwd 必须显示项目根目录。若不是，请先执行：")
    add_code(doc, "cd /jzk-model/jzh_model/bee/bee-paper3-colony-strength-reproduction")
    add_note(doc, "Windows 用户", "prepare_dataset.py 会创建符号链接。推荐使用 VS Code + WSL 或 Remote SSH 在 Linux 环境运行。Windows 原生环境可能需要开启开发者模式或管理员权限才能创建符号链接。")

    doc.add_heading("3. 完整目录结构", level=1)
    doc.add_paragraph("下面的树形结构同时列出源码、数据、权重、训练输出和文档。标记“本地生成/不上传”的目录通常被 .gitignore 排除。")
    add_code(doc, """bee-paper3-colony-strength-reproduction/
├── README.md                         # 项目主说明与快速操作
├── PAPER_NOTES.md                    # 从论文提取的方法和参数
├── REPORT.md                         # 本轮复现结果的 Markdown 报告
├── requirements.txt                  # Python 依赖版本
├── setup.sh                          # 自动创建 .venv 并安装依赖
├── .gitignore                        # Git 忽略数据、权重和运行输出
├── inspect_archive.py                # 安全检查官方 ZIP
├── prepare_dataset.py                # LabelMe → YOLO，固定 60/20/20 划分
├── train_yolov10.py                  # 训练与断点续训入口
├── evaluate.py                       # 固定测试集评估入口
├── predict.py                        # 图片/目录/视频推理入口
├── generate_report.py                # 生成中文 PDF 结果报告
├── generate_word_manual.py           # 生成本 Word 手册
├── generate_paper_analysis.py        # 生成论文深度分析 Word 文档
├── archive_manifest.json             # ZIP 审计结果（本地生成）
├── dataset.yaml                      # Ultralytics 数据集配置（本地生成）
├── dataset_manifest.json             # 每张图片的固定划分清单（本地生成）
├── downloads/                        # 官方下载文件（不上传）
│   └── supplement.zip                # 约 4.2 GB 官方附件
├── data/                             # 解压与转换数据（不上传）
│   ├── raw/                          # 官方 ZIP 解压内容
│   └── yolo/
│       ├── images/{train,val,test}/  # 图片符号链接
│       └── labels/{train,val,test}/  # YOLO TXT 标签
├── weights/                          # 初始预训练权重（不上传）
│   ├── yolov10x.pt                   # 本项目默认初始权重
│   └── yolo26n.pt                    # 其他测试权重，默认训练不使用
├── runs/                             # 训练、评估、推理运行产物（不上传）
│   ├── detector/yolov10x_seed50/
│   │   ├── args.yaml                 # 本次训练实际参数
│   │   ├── results.csv               # 每轮损失与指标
│   │   ├── results.png               # 训练曲线
│   │   ├── confusion_matrix*.png     # 验证集混淆矩阵
│   │   └── weights/{best.pt,last.pt} # 最佳/最近 checkpoint
│   ├── evaluation/test/              # 测试集曲线、混淆矩阵、预测图
│   └── predict/demo/                 # 自定义媒体推理结果
├── results/                          # 可提交的小型最终指标
│   ├── detector_metrics.json
│   ├── experiment_summary.json
│   └── training_metrics.csv
├── reports/                          # PDF 结果报告
├── docs/                             # Word 操作手册
├── .venv/                            # Python 虚拟环境（不上传）
└── __pycache__/                      # Python 字节码缓存（可删除）""")

    doc.add_heading("4. 根目录源码和配置文件逐项说明", level=1)
    doc.add_heading("4.1 主要 Python 脚本", level=2)
    add_table(doc, ["文件", "输入", "输出", "什么时候使用"], [
        ["inspect_archive.py", "downloads/supplement.zip", "archive_manifest.json", "首次下载后检查文件大小、SHA-256、ZIP 内部结构和危险路径"],
        ["prepare_dataset.py", "官方 ZIP 内的 LabelMe JSON 和图片", "data/yolo、dataset.yaml、dataset_manifest.json", "第一次准备数据，或改变随机种子、类别映射后重建数据"],
        ["train_yolov10.py", "dataset.yaml、初始权重或 last.pt", "runs/detector/<实验名>/", "开始新训练或从 checkpoint 续训"],
        ["evaluate.py", "best.pt、dataset.yaml 的 test 划分", "results/detector_metrics.json、测试曲线和混淆矩阵", "训练完成后做独立测试"],
        ["predict.py", "best.pt、图片/目录/视频", "runs/predict/demo/、终端类别计数", "对自己的媒体进行推理"],
        ["generate_report.py", "训练 CSV、测试 JSON、曲线图", "reports/第三篇论文复现结果报告.pdf", "评估完成后生成论文复现报告"],
        ["generate_word_manual.py", "项目结构和固定说明", "docs/…操作手册.docx", "重新生成本 Word 手册"],
        ["generate_paper_analysis.py", "论文全文、论文表格与复现结果", "docs/…深度分析与创新改进建议.docx", "生成通俗论文解读、批判性分析和创新路线"],
    ], [3.2, 4.5, 4.8, 5.0])
    doc.add_heading("4.2 文档与环境文件", level=2)
    add_table(doc, ["文件", "作用", "修改建议"], [
        ["README.md", "项目首页、核心命令、最终指标、文件说明", "操作方法变化后同步修改"],
        ["PAPER_NOTES.md", "论文数据规模、超参数、增强方式和复现限制", "只记录可从论文确认的内容"],
        ["REPORT.md", "纯文本版本的最终复现结果", "新实验完成后更新指标和结论"],
        ["requirements.txt", "锁定 Ultralytics 等 Python 依赖范围", "增加新库时追加，避免无版本约束"],
        ["setup.sh", "创建 .venv、升级 pip、安装 requirements", "通常不需要修改"],
        [".gitignore", "阻止数据、权重、虚拟环境和 runs 进入 Git", "除非使用 Git LFS，否则不要取消大文件忽略"],
    ], [3.5, 7.0, 6.0])
    doc.add_heading("4.3 本地生成配置", level=2)
    add_table(doc, ["文件", "含义", "注意事项"], [
        ["dataset.yaml", "告诉 Ultralytics 数据根目录、train/val/test 相对路径和类别名称", "路径或类别变化后必须更新；prepare_dataset.py 会重写它"],
        ["dataset_manifest.json", "记录 4,590 个样本的来源、输出名、划分和框数", "用于证明本次划分可复查；文件较大但很重要"],
        ["archive_manifest.json", "记录 ZIP 字节数、哈希、文件类型和示例路径", "确认自己下载的附件是否与复现实验一致"],
    ], [4.0, 7.0, 5.5])

    doc.add_heading("5. 数据目录与标签格式", level=1)
    doc.add_heading("5.1 原始数据和转换数据的区别", level=2)
    add_bullets(doc, [
        "downloads/supplement.zip 是原始下载包，不应直接修改。",
        "data/raw/ 是解压结果，保留原始 LabelMe JSON 与图片结构。",
        "data/yolo/ 是训练使用的数据。images 下是原图符号链接，labels 下是 YOLO 文本标签。",
        "删除 data/yolo 后可以重新转换；不要为了修改划分而手工搬动少量文件，应通过脚本和 seed 重建。",
    ])
    doc.add_heading("5.2 YOLO 标签每一行代表什么", level=2)
    add_code(doc, "class_id  center_x  center_y  width  height\n0 0.51234567 0.43876543 0.08200000 0.11600000")
    doc.add_paragraph("后四个数字都按图片宽高归一化到 0～1。class_id 必须从 0 开始连续编号。本项目的 0、1、2 分别对应无花粉工蜂、带花粉工蜂和雄蜂。")
    doc.add_heading("5.3 dataset.yaml 的含义", level=2)
    add_code(doc, """path: /项目绝对路径/data/yolo
train: images/train
val: images/val
test: images/test
names:
  0: WB no pollen
  1: WB pollen
  2: Drone bee""")
    add_note(doc, "类别顺序", "dataset.yaml、prepare_dataset.py 的 CLASS_NAMES/normalized_label、evaluate.py 的 NAMES 必须保持一致。改变类别后应开始新实验，不能继续使用旧 checkpoint。")

    doc.add_heading("6. 第一次在 VS Code 中准备环境", level=1)
    doc.add_heading("6.1 自动安装", level=2)
    add_code(doc, "chmod +x setup.sh\n./setup.sh")
    doc.add_paragraph("setup.sh 会在项目根目录创建 .venv，并按照 requirements.txt 安装依赖。执行完成后，不要在训练命令中使用系统 python，应明确使用 .venv/bin/python。")
    doc.add_heading("6.2 检查 Python、PyTorch 与 GPU", level=2)
    add_code(doc, ".venv/bin/python -c \"import torch; print('torch=', torch.__version__); print('cuda=', torch.cuda.is_available()); print('gpu=', torch.cuda.get_device_name(0) if torch.cuda.is_available() else 'CPU')\"\nnvidia-smi")
    add_table(doc, ["检查结果", "含义", "处理"], [
        ["cuda=True 且显示 GPU 名称", "可以进行 GPU 训练", "继续准备数据"],
        ["cuda=False", "当前 PyTorch 看不到 CUDA", "检查驱动、容器 GPU 映射和 PyTorch CUDA 版本"],
        ["ModuleNotFoundError", "VS Code 或终端使用了错误解释器", "重新选择 .venv/bin/python 或执行 setup.sh"],
        ["nvidia-smi 不存在/失败", "驱动或远程环境没有暴露 GPU", "先修复系统环境，不要直接开始训练"],
    ], [4.0, 6.0, 6.5])

    doc.add_heading("7. 从零准备官方数据", level=1)
    doc.add_heading("7.1 放置压缩包", level=2)
    add_code(doc, "mkdir -p downloads\n# 将官方文件保存为：\n# downloads/supplement.zip\nls -lh downloads/supplement.zip")
    doc.add_paragraph("本次复现的官方 ZIP 大小为 4,236,622,867 字节，SHA-256 为：")
    add_code(doc, "cae66fc3d855ef1d26f570afe574cbfe23763c5da845b100625d0d490fd6c891")
    doc.add_heading("7.2 审计压缩包", level=2)
    add_code(doc, ".venv/bin/python inspect_archive.py")
    doc.add_paragraph("该步骤会检查危险路径和符号链接，并生成 archive_manifest.json。若哈希不同，应先确认是否下载了错误版本或文件损坏。")
    doc.add_heading("7.3 转换并划分", level=2)
    add_code(doc, ".venv/bin/python prepare_dataset.py")
    doc.add_paragraph("如果 data/yolo 已存在，脚本会拒绝覆盖。确认需要重建后使用：")
    add_code(doc, ".venv/bin/python prepare_dataset.py --force")
    add_note(doc, "--force 的影响", "它会删除并重建 data/yolo，但不会删除 downloads/supplement.zip。重建前建议保存旧 dataset_manifest.json，以便比较划分。")
    doc.add_heading("7.4 验证数据是否完整", level=2)
    add_code(doc, "find data/yolo/images/train -type l | wc -l\nfind data/yolo/images/val   -type l | wc -l\nfind data/yolo/images/test  -type l | wc -l")
    add_table(doc, ["划分", "应有图片数", "原始框数"], [
        ["train", "2,754", "47,252"],
        ["val", "918", "15,448"],
        ["test", "918", "16,112"],
    ], [5.0, 5.0, 5.0])

    doc.add_heading("8. 手动开始一次新训练", level=1)
    doc.add_heading("8.1 论文参数的默认训练", level=2)
    add_code(doc, ".venv/bin/python train_yolov10.py")
    doc.add_paragraph("该命令使用 weights/yolov10x.pt、dataset.yaml、batch=4、imgsz=1080、SGD、lr0=0.01、seed=50，并把结果写入 runs/detector/yolov10x_seed50/。")
    add_note(doc, "避免覆盖", "脚本设置 exist_ok=True。重复使用相同 --name 可能复用同一目录，给结果管理带来混乱。新实验务必使用新的 --name。")
    doc.add_heading("8.2 建议使用明确参数启动", level=2)
    add_code(doc, """.venv/bin/python train_yolov10.py \\
  --model weights/yolov10x.pt \\
  --data dataset.yaml \\
  --epochs 100 \\
  --batch 4 \\
  --imgsz 1080 \\
  --device 0 \\
  --workers 8 \\
  --seed 50 \\
  --name yolov10x_seed50_manual""")
    doc.add_heading("8.3 命令行参数说明", level=2)
    add_table(doc, ["参数", "默认值", "作用", "修改建议"], [
        ["--model", "weights/yolov10x.pt", "新训练使用的初始模型/权重", "换模型时指定另一个本地 .pt；不要用于 --resume"],
        ["--data", "dataset.yaml", "数据配置文件", "使用新数据时指向新的 YAML"],
        ["--epochs", "100", "最大训练轮数", "有 early stopping，不代表一定跑满"],
        ["--batch", "4", "每次送入 GPU 的图片数", "显存不足先降低为 2 或 1"],
        ["--imgsz", "1080", "请求输入尺寸", "Ultralytics 会补齐到步长 32 的倍数，本项目实际为 1088"],
        ["--device", "0", "GPU 编号", "第二张卡可用 1；CPU 使用 cpu"],
        ["--workers", "8", "数据加载进程数", "内存不足或 Windows 问题时降为 2/0"],
        ["--seed", "50", "随机种子", "比较实验时固定；研究稳定性时使用多组 seed"],
        ["--name", "yolov10x_seed50", "runs 下实验目录名", "每次新实验必须唯一且能表达参数"],
        ["--resume", "关闭", "从对应实验的 last.pt 完整恢复", "只用于中断恢复，不用于更换模型/类别"],
    ], [2.5, 3.0, 5.5, 6.0])
    doc.add_heading("8.4 在 VS Code 中监控训练", level=2)
    add_bullets(doc, [
        "终端会显示 epoch、GPU 显存、box/cls/dfl loss、Precision、Recall、mAP。",
        "另开一个 VS Code 终端执行 nvidia-smi，观察显存与 GPU 利用率。",
        "results.csv 每个完整 epoch 增加一行；只有写入完整行后，才表示该轮完成。",
        "best.pt 只在验证 fitness 改善时更新；last.pt 每轮更新。",
        "需要停止时在训练终端按 Ctrl+C，尽量在一轮完成后停止。",
    ])
    add_code(doc, "tail -n 3 runs/detector/yolov10x_seed50_manual/results.csv\nwatch -n 2 nvidia-smi")

    doc.add_heading("9. 权重文件的区别与正确用法", level=1)
    add_table(doc, ["权重", "位置", "内容", "使用方式"], [
        ["初始预训练权重", "weights/yolov10x.pt", "COCO 预训练模型，用于开始新实验", "--model weights/yolov10x.pt"],
        ["最佳权重", "runs/detector/<name>/weights/best.pt", "验证 fitness 最好的 epoch，已去除优化器状态", "评估和部署优先使用"],
        ["最近权重", "runs/detector/<name>/weights/last.pt", "最近完整 epoch，训练结束后通常也会精简", "训练中断时由 --resume 加载"],
        ["其他测试权重", "weights/yolo26n.pt", "当前目录中的其他小模型权重", "默认流程不使用，除非明确做新实验"],
    ], [3.0, 5.2, 5.4, 3.7])
    doc.add_heading("9.1 用自己的权重开始新实验", level=2)
    add_code(doc, ".venv/bin/python train_yolov10.py --model /绝对路径/your_model.pt --name custom_pretrained_v1")
    doc.add_paragraph("如果自己的权重类别数与本项目不同，Ultralytics 通常会重建检测头并迁移可匹配层。必须检查启动日志中的 transferred 参数数量和类别数，不能只看程序是否启动。")
    doc.add_heading("9.2 用最佳权重做评估或推理", level=2)
    add_code(doc, ".venv/bin/python evaluate.py --weights runs/detector/实验名/weights/best.pt\n.venv/bin/python predict.py your.jpg --weights runs/detector/实验名/weights/best.pt")
    add_note(doc, "不要混用", "开始新实验可以加载 best.pt 做微调，但这不等同于 --resume。--resume 恢复原实验状态；加载 best.pt 开新实验会创建新的优化器和 epoch 计数。")

    doc.add_heading("10. 断点续训：关机后如何继续", level=1)
    doc.add_heading("10.1 当前默认实验续训", level=2)
    add_code(doc, ".venv/bin/python train_yolov10.py --resume")
    doc.add_heading("10.2 指定其他实验目录续训", level=2)
    add_code(doc, ".venv/bin/python train_yolov10.py --resume --name yolov10x_seed50_manual")
    doc.add_paragraph("脚本会拼接 runs/detector/<name>/weights/last.pt，并调用 model.train(resume=True)。Ultralytics 从 checkpoint 恢复 epoch、优化器、学习率调度和 early stopping 状态。")
    doc.add_heading("10.3 续训前的检查", level=2)
    add_code(doc, "ls -lh runs/detector/yolov10x_seed50_manual/weights/last.pt\ntail -n 2 runs/detector/yolov10x_seed50_manual/results.csv")
    add_bullets(doc, [
        "如果 last.pt 不存在，不能使用 --resume。",
        "续训时 --model 通常不会生效，因为模型来源固定为 last.pt。",
        "不要在改变类别数、dataset.yaml 或网络结构后继续旧 checkpoint，应开新实验。",
        "仅改变训练参数时，也应先确认 Ultralytics 的 resume 会恢复 checkpoint 中的旧参数；需要新参数时更适合加载权重开新实验。",
    ])

    doc.add_heading("11. 如何调整训练参数", level=1)
    doc.add_heading("11.1 不改代码即可调整的参数", level=2)
    add_code(doc, "# 12 GB 左右显存的保守示例\n.venv/bin/python train_yolov10.py --batch 2 --imgsz 960 --workers 4 --name x_b2_i960\n\n# 快速验证流程是否能跑通\n.venv/bin/python train_yolov10.py --epochs 2 --batch 1 --imgsz 640 --name smoke_test")
    doc.add_heading("11.2 当前写死在 train_yolov10.py 中的参数", level=2)
    add_table(doc, ["代码参数", "当前值", "含义", "常见调整"], [
        ["patience", "10", "连续多少轮无改善后早停", "探索长尾收敛可改 20；固定轮数可改 0 关闭早停"],
        ["optimizer", "SGD", "优化器", "对照实验可尝试 AdamW，但不再严格匹配论文"],
        ["lr0", "0.01", "初始学习率", "batch 很小时可尝试更低学习率"],
        ["lrf", "0.01", "最终学习率比例", "配合调度策略修改"],
        ["cos_lr", "False", "是否使用余弦学习率", "启用后形成新实验"],
        ["fliplr", "0.4", "水平翻转概率", "场景方向敏感时降低"],
        ["degrees", "10.0", "随机旋转角度范围", "摄像机固定时不宜过大"],
        ["deterministic", "True", "优先可重复性", "关闭可能更快，但结果更难严格重复"],
    ], [3.0, 2.3, 6.0, 6.0])
    doc.add_paragraph("修改方式：在 VS Code 打开 train_yolov10.py，找到 model.train(...) 调用，修改对应关键字参数，然后使用新的 --name 开始训练。")
    add_code(doc, """model.train(
    ...
    patience=20,
    optimizer="AdamW",
    lr0=0.001,
    cos_lr=True,
    ...
)""")
    add_note(doc, "实验可比性", "一次只改一到两个关键变量，并记录实验名、Git 提交、args.yaml、数据 manifest 和最终指标。否则无法判断提升来自哪个改动。")

    doc.add_heading("12. 如何更换模型大小或结构", level=1)
    doc.add_paragraph("当前论文复现使用 YOLOv10X。若只为学习和快速试验，可以更换更小权重；若用于论文严格对照，应保留 YOLOv10X。")
    add_table(doc, ["选择", "显存/速度趋势", "精度趋势", "适用场景"], [
        ["YOLOv10N/S", "显存小、速度快", "通常低于 X", "流程验证、低功耗部署"],
        ["YOLOv10M/B/L", "居中", "居中", "算力与精度折中"],
        ["YOLOv10X", "显存最大、训练慢", "通常最高", "论文复现与服务器部署"],
    ], [3.5, 4.5, 4.0, 4.5])
    doc.add_paragraph("先把权重文件放入 weights/，再指定独立实验名：")
    add_code(doc, ".venv/bin/python train_yolov10.py --model weights/yolov10s.pt --batch 8 --name yolov10s_seed50")
    doc.add_paragraph("模型文件名必须真实存在。若 Ultralytics 自动尝试联网下载，而服务器没有网络，命令会失败；因此推荐提前下载权重并使用本地路径。")

    doc.add_heading("13. 如何替换数据集或增加类别", level=1)
    doc.add_heading("13.1 使用已经是 YOLO 格式的新数据", level=2)
    add_steps(doc, [
        "建立新的 images/train、images/val、images/test 与对应 labels 目录。",
        "确认每张图片有同名 .txt 标签；无目标图片可以对应空标签。",
        "创建新的 YAML，填写 path、train、val、test 和 names。",
        "用 --data 指向新 YAML，并使用新的 --name 开始训练。",
    ])
    add_code(doc, ".venv/bin/python train_yolov10.py --data configs/my_dataset.yaml --name my_dataset_v1")
    doc.add_heading("13.2 修改本项目的 LabelMe 类别映射", level=2)
    add_bullets(doc, [
        "在 prepare_dataset.py 修改 CLASS_NAMES。",
        "在 normalized_label() 中明确每个 LabelMe 标签如何映射到 class_id。",
        "在 evaluate.py 修改 NAMES，使逐类别报告顺序一致。",
        "删除或备份旧 data/yolo，使用 --force 重新转换。",
        "检查 dataset.yaml 中类别编号从 0 连续排列。",
        "使用新的实验名从初始权重开始训练。",
    ])
    add_note(doc, "防止静默错误", "最危险的问题是标签 ID 与 names 顺序不一致。程序可能正常训练，但类别含义全部错位。转换后应随机打开多张图片和标签做可视化核对。")

    doc.add_heading("14. 测试集评估与指标解读", level=1)
    doc.add_heading("14.1 运行评估", level=2)
    add_code(doc, """.venv/bin/python evaluate.py \\
  --weights runs/detector/实验名/weights/best.pt \\
  --data dataset.yaml \\
  --imgsz 1080 \\
  --batch 4 \\
  --device 0""")
    doc.add_heading("14.2 评估输出", level=2)
    add_bullets(doc, [
        "results/detector_metrics.json：整体和逐类别 Precision、Recall、F1、AP50、mAP@0.5:0.95、速度和权重 SHA-256。",
        "runs/evaluation/test/confusion_matrix_normalized.png：归一化混淆矩阵。",
        "runs/evaluation/test/BoxPR_curve.png：Precision-Recall 曲线。",
        "val_batch*_labels.jpg / pred.jpg：真实标注与预测对照。",
    ])
    doc.add_heading("14.3 指标如何判断", level=2)
    add_table(doc, ["指标", "说明", "看高还是看低"], [
        ["Precision", "预测出的框有多少是正确的", "越高越好；低表示误检多"],
        ["Recall", "真实目标有多少被检出", "越高越好；低表示漏检多"],
        ["F1", "Precision 与 Recall 的调和平均", "越高越均衡"],
        ["AP50", "IoU=0.5 条件下的平均精度", "越高越好，定位要求相对宽松"],
        ["mAP@0.5:0.95", "多个 IoU 阈值的严格综合指标", "越高越好，更能反映定位质量"],
        ["box/cls/dfl loss", "训练和验证损失", "一般下降；需结合 mAP 判断，不能单看 loss"],
    ], [3.5, 8.0, 5.0])

    doc.add_heading("15. 对图片、目录和视频进行推理", level=1)
    doc.add_heading("15.1 单张图片", level=2)
    add_code(doc, ".venv/bin/python predict.py /path/to/bee.jpg --weights runs/detector/实验名/weights/best.pt --conf 0.25")
    doc.add_heading("15.2 整个图片目录", level=2)
    add_code(doc, ".venv/bin/python predict.py /path/to/images --weights runs/detector/实验名/weights/best.pt")
    doc.add_heading("15.3 视频", level=2)
    add_code(doc, ".venv/bin/python predict.py /path/to/video.mp4 --weights runs/detector/实验名/weights/best.pt --conf 0.30")
    add_table(doc, ["参数", "影响", "建议"], [
        ["--conf 降低", "检出更多，同时可能增加误检", "漏检明显时从 0.25 降到 0.15～0.20 试验"],
        ["--conf 提高", "结果更保守，误检减少，可能漏检", "部署前按验证集和业务成本选择"],
        ["--imgsz 提高", "小目标可能更清楚，但显存与耗时增加", "保持训练尺寸附近，避免无依据放大"],
    ], [3.0, 7.0, 6.5])
    add_note(doc, "计数含义", "predict.py 对视频输出的是每帧检测框累计数，同一只蜜蜂跨帧会重复计数。如果需要个体进出数量，必须增加多目标跟踪和过线计数。")

    doc.add_heading("16. runs 目录中的训练产物怎么看", level=1)
    add_table(doc, ["产物", "如何使用", "是否需要长期保留"], [
        ["args.yaml", "记录实际运行参数，是复现实验的首要依据", "保留"],
        ["results.csv", "每轮损失、P/R、mAP 和学习率，可用于画图和比较", "保留"],
        ["results.png", "训练/验证损失与指标随 epoch 的曲线", "保留"],
        ["best.pt", "评估、推理、部署", "必须保留或备份"],
        ["last.pt", "训练中断恢复", "训练完成后按需要保留"],
        ["train_batch*.jpg", "检查增强后的训练图片和标签是否合理", "调试期保留"],
        ["val_batch*_labels/pred.jpg", "对照真实框与预测框", "误差分析时保留"],
        ["train_resume.log", "排查续训、错误和完整输出", "有问题时保留，可能较大"],
    ], [4.0, 8.0, 4.5])

    doc.add_heading("17. 推荐的实验管理方法", level=1)
    doc.add_heading("17.1 实验命名", level=2)
    add_code(doc, "模型_数据版本_关键参数_种子\n例如：yolov10x_bee_v1_b4_i1080_s50\n例如：yolov10s_bee_v1_b8_i960_s50")
    doc.add_heading("17.2 每次实验至少保存", level=2)
    add_bullets(doc, [
        "实验对应的代码 Git commit。",
        "dataset_manifest.json 或数据版本哈希。",
        "runs/<实验>/args.yaml。",
        "results.csv、best.pt 和最终测试 JSON。",
        "硬件、CUDA、PyTorch、Ultralytics 版本。",
        "与基线相比只改了哪些变量。",
    ])
    doc.add_heading("17.3 使用 Git 保护代码修改", level=2)
    add_code(doc, "git status\ngit diff\ngit add train_yolov10.py README.md\ngit commit -m \"Tune learning rate for bee detector\"")
    doc.add_paragraph("数据集、runs 和 .pt 权重默认不会提交。若确实要在远端保存大权重，应使用 Git LFS、Release、对象存储或模型仓库，避免直接写入普通 Git 历史。")

    doc.add_heading("18. 常见错误与处理", level=1)
    add_table(doc, ["现象/错误", "常见原因", "处理方法"], [
        ["CUDA out of memory", "batch/imgsz 太大，或 GPU 上有其他进程", "先减 batch，再减 imgsz；用 nvidia-smi 清查进程"],
        ["No module named ultralytics", "解释器选错或依赖未安装", "使用 .venv/bin/python，重新执行 setup.sh"],
        ["No such file: dataset.yaml", "未运行数据转换或终端目录错误", "运行 prepare_dataset.py，并检查 pwd"],
        ["Output exists; use --force", "data/yolo 已存在", "确认后使用 prepare_dataset.py --force，或保留现有划分"],
        ["last.pt 不存在", "实验名写错、未完成首轮或 runs 被删除", "检查 --name 和 weights 目录；无法恢复时开新实验"],
        ["imgsz 更新为 1088", "1080 不是最大步长 32 的倍数", "这是正常自动补齐，不是错误"],
        ["训练突然很慢", "GPU 未使用、数据 I/O、显存换页或其他任务占用", "检查 nvidia-smi、CUDA、workers 和磁盘"],
        ["DataLoader worker exited", "内存不足、Windows 多进程或坏图片", "workers 降为 0/2，定位损坏样本"],
        ["预测没有框", "权重/类别错误、conf 太高或输入域差异大", "确认 best.pt，降低 conf，检查训练类别和样例图"],
        ["评估实例数少 1", "官方数据含一个完全重复框，Ultralytics 自动去重", "本项目测试按 16,111 个有效实例计分，属于已知现象"],
        ["GitHub 克隆后没有数据/权重", ".gitignore 排除了大文件", "按 README 从官方来源下载数据，并重新训练或复制权重"],
        ["Word 目录没有页码", "目录域尚未更新", "右键目录 → 更新域 → 更新整个目录"],
    ], [4.2, 5.2, 7.0])

    doc.add_heading("19. 不同硬件的训练建议", level=1)
    add_table(doc, ["硬件", "建议起点", "说明"], [
        ["RTX 4090 24 GB", "YOLOv10X、batch 4、imgsz 1080", "本次复现使用配置"],
        ["16 GB GPU", "YOLOv10X batch 2，或更小模型 batch 4", "根据实际显存逐步提高"],
        ["8～12 GB GPU", "较小模型、batch 1～2、imgsz 640～960", "先做 smoke test，再决定精度/速度折中"],
        ["CPU", "只做少量推理或流程测试", "YOLOv10X 完整训练耗时通常不可接受"],
    ], [4.0, 7.0, 5.5])
    doc.add_heading("19.1 先做两轮冒烟测试", level=2)
    add_code(doc, ".venv/bin/python train_yolov10.py --epochs 2 --batch 1 --imgsz 640 --name smoke_test")
    doc.add_paragraph("冒烟测试只检查环境、数据、标签和权重能否正常运行，不用于报告最终精度。确认无误后使用正式参数和新实验名重新训练。")

    doc.add_heading("20. 当前复现实验结果示例", level=1)
    add_table(doc, ["项目", "结果"], [
        ["自然停止轮次", "43（patience=10）"],
        ["最佳轮次", "33"],
        ["测试 Precision / Recall / F1", "0.9455 / 0.9254 / 0.9353"],
        ["测试 mAP@0.5", "0.9629"],
        ["测试 mAP@0.5:0.95", "0.6374"],
        ["最佳权重 SHA-256", "46e72c5b4eb13251a160a5a932616196226f18546f6116f0bea1cc9be4fbc407"],
    ], [5.5, 11.0])
    figures = [
        (ROOT / "runs/detector/yolov10x_seed50/results.png", "图 1  当前复现实验训练曲线"),
        (ROOT / "runs/evaluation/test/confusion_matrix_normalized.png", "图 2  固定测试集归一化混淆矩阵"),
        (ROOT / "runs/predict/demo/GH010942.MP4_frame_1005_72780290a2.jpg", "图 3  单张测试图片推理示例"),
    ]
    for image_path, caption in figures:
        if image_path.exists():
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Inches(6.2))
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)
    add_note(doc, "结果比较", "论文未公开逐帧划分和完整 accuracy/kappa 实现。本项目标准检测指标可在自身固定划分上重复，但不应把 mAP 直接改名为论文 accuracy。")

    doc.add_heading("21. 一页命令速查", level=1)
    add_code(doc, """# 进入项目
cd /jzk-model/jzh_model/bee/bee-paper3-colony-strength-reproduction

# 安装环境
chmod +x setup.sh && ./setup.sh

# 检查 GPU
.venv/bin/python -c "import torch; print(torch.cuda.is_available(), torch.cuda.get_device_name(0))"

# 审计与准备数据
.venv/bin/python inspect_archive.py
.venv/bin/python prepare_dataset.py

# 新训练（务必使用唯一实验名）
.venv/bin/python train_yolov10.py --name my_exp_v1

# 中断后续训
.venv/bin/python train_yolov10.py --resume --name my_exp_v1

# 评估最佳权重
.venv/bin/python evaluate.py --weights runs/detector/my_exp_v1/weights/best.pt

# 图片/视频推理
.venv/bin/python predict.py /path/to/input --weights runs/detector/my_exp_v1/weights/best.pt

# 生成 PDF 与 Word 文档
.venv/bin/python generate_report.py
.venv/bin/python generate_word_manual.py
.venv/bin/python generate_paper_analysis.py

# 查看训练进度
tail -n 3 runs/detector/my_exp_v1/results.csv
nvidia-smi""")

    doc.add_heading("22. 自己修改项目时的推荐流程", level=1)
    add_steps(doc, [
        "先完整跑通 smoke_test，确认当前代码和数据没有问题。",
        "用 git status 和 git diff 记录修改前状态，必要时创建分支。",
        "只修改一个明确目标，例如 batch、学习率、模型大小或数据增强。",
        "给新实验设置唯一 --name，不覆盖基线目录。",
        "训练结束后始终用固定 test 划分评估 best.pt。",
        "对比整体与逐类别指标，尤其关注样本最少的 Drone bee。",
        "记录 args.yaml、results.csv、权重哈希和代码提交。",
        "确认改动有效后，再更新 README、REPORT 和报告文档。",
    ])
    add_note(doc, "最重要原则", "不要同时修改数据划分、模型、学习率和增强后只比较一个 mAP。控制变量、保存实验记录，才能真正理解代码和训练结果。")

    doc.add_heading("附录 A：需要备份与可以删除的内容", level=1)
    add_table(doc, ["类型", "内容", "建议"], [
        ["必须备份", "best.pt、args.yaml、results.csv、测试 JSON、代码 commit", "实验可复现的核心"],
        ["建议备份", "last.pt、曲线、混淆矩阵、关键预测样例", "便于续训和误差分析"],
        ["可重新生成", "data/yolo、dataset.yaml、manifest、报告文件", "前提是官方 ZIP 与源码仍在"],
        ["通常可删除", "__pycache__、临时日志、失败实验的大量预测图片", "删除前确认不再需要排错"],
        ["不要提交普通 Git", "supplement.zip、data、.venv、runs、.pt", "使用外部存储或 Git LFS"],
    ], [3.3, 7.5, 5.7])

    doc.add_heading("附录 B：手册覆盖范围", level=1)
    doc.add_paragraph("本手册针对当前第三篇论文 YOLOv10X 项目和现有脚本编写。它解释了当前代码真实支持的命令行参数，并把需要修改源码的参数明确区分出来。若以后重构训练框架、加入 Faster R-CNN/RetinaNet/FCOS/SSD、增加多目标跟踪或更换数据格式，应同步更新本手册。")
    doc.add_paragraph("项目主说明：README.md\n论文方法摘录：PAPER_NOTES.md\n结果报告：REPORT.md 与 reports/第三篇论文复现结果报告.pdf")
    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_manual()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
